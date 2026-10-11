import math
import os
import pickle
import re
from pathlib import Path

import numpy as np
import requests
from openpyxl import Workbook, load_workbook


# ------------------------------------------------------------
# 参数定义
# ------------------------------------------------------------

ROOT_PATH = Path(__file__).resolve().parent

ENV_FILE = ".env"
ENV_PATH = ROOT_PATH / ENV_FILE


INPUT_FILE = "clean_output.xlsx"
INPUT_PATH = ROOT_PATH / INPUT_FILE

OUTPUT_FILE = "clean_output_labeled.xlsx"
OUTPUT_PATH = ROOT_PATH / OUTPUT_FILE

CACHE_FILE = "location_embedding_cache.pkl"
CACHE_PATH = ROOT_PATH / CACHE_FILE

OUTPUT_SHEET = "labeled"
SUMMARY_SHEET = "clusters"

API_URL = os.environ['EMBEDDING_API_URL']
API_KEY_ENV = "SILICONFLOW_API_KEY"
MODEL = "Qwen/Qwen3-Embedding-8B"
DIMENSIONS = 1536
BATCH_SIZE = 64

W_CORE = 0.5
W_POI = 0.3
W_ROAD_NO = 0.20

K_ROAD_NO = 150
THRESHOLD_LINK = 0.8

LOCATION_FIELDS = ["district", "road_name", "road_no", "entity_name"]
DETAIL_FIELDS = [
    "location_label",
]
SUMMARY_FIELDS = [
    "location_label",
    "district",
    "road_name",
    "road_no",
    "entity_name",
    "node_count",
    "company_count",
    "sample_companies",
]


# ------------------------------------------------------------
# 工具函数
# ------------------------------------------------------------

def text(v):
    if v is None:
        return ""
    return str(v).strip()


def load_env_file(path):
    if not path.exists():
        return

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("export "):
            line = line[len("export ") :].strip()
        if "=" not in line:
            continue

        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip()
        if not key:
            continue
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
            value = value[1:-1]

        os.environ.setdefault(key, value)


def load_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = list(next(rows, []) or [])
    data = [list(row) for row in rows]
    workbook.close()
    return header, data


def build_nodes(rows, idx):
    nodes = []
    node_map = {}
    for row_id, row in enumerate(rows):
        key = tuple(text(row[idx[name]] if idx[name] < len(row) else None) for name in LOCATION_FIELDS)
        node_id = node_map.get(key)
        if node_id is None:
            node_id = len(nodes)
            node_map[key] = node_id
            node = dict(zip(LOCATION_FIELDS, key))
            node["row_ids"] = []
            node["count"] = 0
            node["sample_companies"] = []
            node["core_text"] = make_core_input(node)
            node["poi_text"] = make_poi_input(node)
            node["road_no_value"] = parse_road_no(node["road_no"])
            nodes.append(node)
        node = nodes[node_id]
        node["row_ids"].append(row_id)
        node["count"] += 1
        company_name = text(row[idx["company_name"]] if idx["company_name"] < len(row) else None)
        if company_name and company_name not in node["sample_companies"] and len(node["sample_companies"]) < 10:
            node["sample_companies"].append(company_name)
    return nodes


def make_core_input(node):
    return " ".join([part for part in [node["district"], node["road_name"]] if part])


def make_poi_input(node):
    return node["entity_name"]


def get_api_key():
    api_key = os.getenv(API_KEY_ENV, "").strip()
    if not api_key:
        raise ValueError(f"缺少 API key，请先设置环境变量 {API_KEY_ENV}")
    return api_key


def load_cache(path):
    if not path.exists():
        return {}
    with path.open("rb") as file:
        data = pickle.load(file)
    if not isinstance(data, dict):
        return {}
    if data.get("model") != MODEL or data.get("dimensions") != DIMENSIONS:
        return {}
    items = data.get("items", {})
    return items if isinstance(items, dict) else {}


def save_cache(path, cache):
    with path.open("wb") as file:
        pickle.dump({"model": MODEL, "dimensions": DIMENSIONS, "items": cache}, file, protocol=pickle.HIGHEST_PROTOCOL)


def fetch_embeddings(texts):
    cache = load_cache(CACHE_PATH)
    wanted = []
    seen = set()
    for item in texts:
        item = text(item)
        if item and item not in seen:
            seen.add(item)
            wanted.append(item)
    missing = [item for item in wanted if item not in cache]
    total = math.ceil(len(missing) / BATCH_SIZE) if missing else 0
    if missing:
        headers = {"Authorization": f"Bearer {get_api_key()}", "Content-Type": "application/json"}
        for start in range(0, len(missing), BATCH_SIZE):
            batch = missing[start : start + BATCH_SIZE]
            batch_no = start // BATCH_SIZE + 1
            print(f"开始请求 embedding，第 {batch_no}/{total} 批 ...")
            response = requests.post(
                API_URL,
                headers=headers,
                json={
                    "model": MODEL,
                    "input": batch,
                    "dimensions": DIMENSIONS,
                    "encoding_format": "float",
                },
                timeout=300,
            )
            if not response.ok:
                raise ValueError(f"embedding 请求失败：HTTP {response.status_code}，响应内容：{response.text}")
            data = response.json().get("data", [])
            if len(data) != len(batch):
                raise ValueError("embedding 返回数量异常")
            for item, embedding in zip(batch, [part["embedding"] for part in sorted(data, key=lambda part: part["index"])]):
                cache[item] = embedding
            save_cache(CACHE_PATH, cache)
    if missing:
        print(f"embedding 缓存已更新：{CACHE_PATH.name}")
    else:
        print(f"embedding 全部命中缓存：{CACHE_PATH.name}")
    return {item: cache[item] for item in wanted}


def cosine(a, b):
    dot = sum(x * y for x, y in zip(a, b))
    norm_a = math.sqrt(sum(x * x for x in a))
    norm_b = math.sqrt(sum(y * y for y in b))
    if norm_a == 0 or norm_b == 0:
        return 0.0
    return dot / (norm_a * norm_b)


def parse_road_no(value):
    nums = [int(item) for item in re.findall(r"\d+", text(value))]
    if not nums:
        return None
    return sum(nums) / len(nums)


def road_no_score(a, b, k):
    if a["road_no_value"] is None or b["road_no_value"] is None:
        return None
    return 1 / (1 + abs(a["road_no_value"] - b["road_no_value"]) / k)


def build_sim_cache(values, embeddings):
    items = [item for item in sorted(set(values)) if item]
    if not items:
        return {}, np.zeros((0, 0), dtype=np.float32)
    matrix = np.asarray([embeddings[item] for item in items], dtype=np.float32)
    norms = np.linalg.norm(matrix, axis=1, keepdims=True)
    norms[norms == 0] = 1.0
    matrix = matrix / norms
    return {item: i for i, item in enumerate(items)}, matrix @ matrix.T


def build_road_no_cache(values, k):
    items = sorted(set(value for value in values if value is not None))
    if not items:
        return {}, np.zeros((0, 0), dtype=np.float32)
    numbers = np.asarray(items, dtype=np.float32)
    return {item: i for i, item in enumerate(items)}, 1.0 / (1.0 + np.abs(numbers[:, None] - numbers[None, :]) / k)


def total_score(a, b, caches):
    scores = [
        (W_CORE, float(caches["core_scores"][a["core_id"], b["core_id"]]) if a["core_id"] is not None and b["core_id"] is not None else None),
        (W_POI, float(caches["poi_scores"][a["poi_id"], b["poi_id"]]) if a["poi_id"] is not None and b["poi_id"] is not None else None),
        (W_ROAD_NO, float(caches["road_scores"][a["road_no_id"], b["road_no_id"]]) if a["road_no_id"] is not None and b["road_no_id"] is not None else None),
    ]
    valid = [(weight, score) for weight, score in scores if score is not None]
    if not valid:
        return 0.0
    return sum(weight * score for weight, score in valid) / sum(weight for weight, _ in valid)


def can_link(a, b, score):
    return score >= THRESHOLD_LINK


def union_find(n):
    parent = list(range(n))
    rank = [0] * n

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a, b):
        ra = find(a)
        rb = find(b)
        if ra == rb:
            return False
        if rank[ra] < rank[rb]:
            parent[ra] = rb
        elif rank[ra] > rank[rb]:
            parent[rb] = ra
        else:
            parent[rb] = ra
            rank[ra] += 1
        return True

    return find, union


def write_result(path, header, rows, labels, clusters):
    workbook = Workbook()
    detail_sheet = workbook.active
    detail_sheet.title = OUTPUT_SHEET
    detail_sheet.append(header + DETAIL_FIELDS)
    for row_id, row in enumerate(rows):
        detail_sheet.append(
            [text(row[i]) if i < len(row) else "" for i in range(len(header))]
            + [
                labels[row_id]["location_label"],
            ]
        )

    summary_sheet = workbook.create_sheet(SUMMARY_SHEET)
    summary_sheet.append(SUMMARY_FIELDS)
    for cluster in clusters:
        summary_sheet.append([cluster[field] for field in SUMMARY_FIELDS])
    workbook.save(path)


# ------------------------------------------------------------
# 核心逻辑
# ------------------------------------------------------------

def run():
    print(f"开始读取输入：{INPUT_PATH.name}")
    header, rows = load_rows(INPUT_PATH)
    if not header:
        raise ValueError("输入文件为空")

    required = ["company_name"] + LOCATION_FIELDS
    missing = [name for name in required if name not in header]
    if missing:
        raise ValueError(f"输入文件缺少列：{', '.join(missing)}")

    idx = {name: header.index(name) for name in required}
    nodes = build_nodes(rows, idx)
    print(f"位置节点去重完成，共 {len(nodes)} 个")

    print(f"开始调用在线模型：{MODEL}")
    core_embeddings = fetch_embeddings([node["core_text"] for node in nodes])
    poi_embeddings = fetch_embeddings([node["poi_text"] for node in nodes if node["poi_text"]])
    core_index, core_scores = build_sim_cache([node["core_text"] for node in nodes], core_embeddings)
    poi_index, poi_scores = build_sim_cache([node["poi_text"] for node in nodes if node["poi_text"]], poi_embeddings)
    road_index, road_scores = build_road_no_cache([node["road_no_value"] for node in nodes], K_ROAD_NO)
    for node in nodes:
        node["core_id"] = core_index.get(node["core_text"])
        node["poi_id"] = poi_index.get(node["poi_text"]) if node["poi_text"] else None
        node["road_no_id"] = road_index.get(node["road_no_value"])
    print("embedding 生成完成")

    print("开始聚类 ...")
    find, union = union_find(len(nodes))
    linked = [[] for _ in nodes]
    caches = {"core_scores": core_scores, "poi_scores": poi_scores, "road_scores": road_scores}
    district_groups = {}
    for node_id, node in enumerate(nodes):
        district_groups.setdefault(node["district"], []).append(node_id)

    for district, node_ids in district_groups.items():
        print(f"处理中：district={district or '(空)'}，节点数={len(node_ids)}")
        for offset, left in enumerate(node_ids):
            for right in node_ids[offset + 1 :]:
                score = total_score(nodes[left], nodes[right], caches)
                if can_link(nodes[left], nodes[right], score):
                    union(left, right)
                    linked[left].append((right, score))
                    linked[right].append((left, score))

    cluster_map = {}
    for node_id in range(len(nodes)):
        cluster_map.setdefault(find(node_id), []).append(node_id)

    ordered_clusters = sorted(cluster_map.values(), key=lambda members: (-len(members), min(members)))
    labels = {}
    node_meta = {}
    clusters = []

    for cluster_no, members in enumerate(ordered_clusters, 1):
        member_set = set(members)
        representative = max(
            members,
            key=lambda node_id: (
                len([other for other, _ in linked[node_id] if other in member_set]),
                sum(score for other, score in linked[node_id] if other in member_set),
                bool(nodes[node_id]["entity_name"]),
                -node_id,
            ),
        )
        representative_node = nodes[representative]
        label = f"LOC_{cluster_no:06d}"
        sample_companies = []
        seen = set()
        for node_id in members:
            for company_name in nodes[node_id]["sample_companies"]:
                if company_name and company_name not in seen:
                    seen.add(company_name)
                    sample_companies.append(company_name)
                if len(sample_companies) >= 10:
                    break
            if len(sample_companies) >= 10:
                break

        for node_id in members:
            node_meta[node_id] = {
                "location_label": label,
            }

        clusters.append(
            {
                "location_label": label,
                "district": representative_node["district"],
                "road_name": representative_node["road_name"],
                "road_no": representative_node["road_no"],
                "entity_name": representative_node["entity_name"],
                "node_count": len(members),
                "company_count": sum(nodes[node_id]["count"] for node_id in members),
                "sample_companies": "、".join(sample_companies),
            }
        )

    for node_id, node in enumerate(nodes):
        for row_id in node["row_ids"]:
            labels[row_id] = node_meta[node_id]

    print(f"开始写出结果：{OUTPUT_PATH.name}")
    write_result(OUTPUT_PATH, header, rows, labels, clusters)
    print("处理完成")


# ------------------------------------------------------------
# 主函数
# ------------------------------------------------------------

def main():
    load_env_file(ENV_PATH)
    run()


if __name__ == "__main__":
    main()
