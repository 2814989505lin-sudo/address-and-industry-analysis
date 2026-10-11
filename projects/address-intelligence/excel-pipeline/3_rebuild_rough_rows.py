import os

import json
import re
import time
import signal
import sys
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook


# ------------------------------------------------------------
# 参数定义
# ------------------------------------------------------------

INPUT_FILE_NAME = "rough_output_input.xlsx"
INPUT_PATH = Path(__file__).with_name(INPUT_FILE_NAME)

OUTPUT_FILE_NAME = "rough_rebuilt_output_input.xlsx"
OUTPUT_PATH = Path(__file__).with_name(OUTPUT_FILE_NAME)

# 硅基流动 API 配置（兼容 OpenAI 格式）
DEEPSEEK_API_URL = os.environ['LLM_API_URL']
DEEPSEEK_API_KEY = os.environ['LLM_API_KEY']
DEEPSEEK_MODEL = "deepseek-ai/DeepSeek-V3"

API_RETRY_TIMES = 5
API_RETRY_SLEEP_SECONDS = 2
SAVE_INTERVAL = 20

STATUS_PENDING = "pending"
STATUS_SUCCESS = "success"
STATUS_FAILED = "failed"

STRUCTURED_FIELDS = [
    "province",
    "city",
    "district",
    "town_street",
    "road_name",
    "road_no",
    "entity_name",
    "Additional Information",
]

# 全局变量用于信号处理
current_header = None
current_rows = None
interrupted = False


# ------------------------------------------------------------
# 工具函数
# ------------------------------------------------------------

def normalize_value(value):
    """把单元格值统一成干净字符串。"""
    if value is None:
        return ""
    return str(value).strip()


def load_rows():
    """读取输入 Excel，要求包含 'company_address' 列。"""
    workbook = load_workbook(INPUT_PATH, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return [], []

    header_list = [normalize_value(h) for h in header]
    # 必须存在 company_address 列
    if "company_address" not in header_list:
        raise ValueError("输入 Excel 缺少 'company_address' 列")

    addr_index = header_list.index("company_address")
    all_rows = []
    for row in rows:
        addr = normalize_value(row[addr_index] if addr_index < len(row) else "")
        all_rows.append({"company_address": addr})
    return header_list, all_rows


def make_row_key(row):
    """用地址生成稳定主键。"""
    return normalize_value(row.get("company_address", ""))


def load_existing_progress():
    """读取已有输出文件，用于断点续跑。"""
    if not OUTPUT_PATH.exists():
        return {}

    workbook = load_workbook(OUTPUT_PATH, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    header = list(next(rows, []) or [])
    if not header:
        workbook.close()
        return {}

    # 找到 company_address 在输出文件中的列位置
    if "company_address" not in header:
        workbook.close()
        return {}

    addr_idx = header.index("company_address")
    progress = {}
    for row in rows:
        if len(row) <= addr_idx:
            continue
        addr = normalize_value(row[addr_idx])
        if not addr:
            continue
        row_dict = {"company_address": addr}
        for idx, field_name in enumerate(header):
            if idx == addr_idx:
                continue
            row_dict[field_name] = normalize_value(row[idx] if idx < len(row) else "")
        progress[addr] = row_dict

    workbook.close()
    return progress


def merge_existing_progress(rows, progress):
    """把旧输出中的成功或失败状态合并到当前输入行。"""
    for row in rows:
        key = make_row_key(row)
        old_row = progress.get(key)
        if not old_row:
            row["run_status"] = STATUS_PENDING
            row["api_error"] = ""
            continue

        for field_name in STRUCTURED_FIELDS:
            row[f"rebuilt_{field_name}"] = normalize_value(old_row.get(field_name, ""))
        row["api_error"] = normalize_value(old_row.get("api_error", ""))
        row["run_status"] = normalize_value(old_row.get("run_status", "")) or STATUS_PENDING


def is_row_successful(row):
    """判断一行是否已经成功处理。"""
    return normalize_value(row.get("run_status", "")) == STATUS_SUCCESS


def build_prompt(row):
    """构造单行重构 prompt。"""
    company_address = normalize_value(row.get("company_address", ""))
    input_payload = {"company_address": company_address}

    return f"""
你是中文地址结构化助手。请只根据输入里的 company_address 做字段抽取。

要求：
1. 所有字段值都必须直接来自 company_address 原文，且必须是原文中的连续子串
2. 不允许补全、猜测、改写或归一化成 company_address 里不存在的写法
3. 同一段文本不要重复出现在多个字段中
4. 只输出一个合法 JSON 对象，不要输出解释
5. 禁止猜测性提取：如果不确定某段文本的归属字段，默认归入 Additional Information

字段框架与定义：
你只能输出以下 8 个字段，且字段名必须完全一致：

- `province`：省级行政区划，如"广东省""北京市"
- `city`：地级市或直辖市，如"广州市""北京市"
- `district`：区、县、县级市，如"白云区""天河区"
- `town_street`：镇、街道、街镇等次级行政区划，如"狮岭镇""石牌街道"
- `road_name`：道路名称，必须是正式的道路名。
- `road_no`：纯路号，如"208号""77-79号""77至79号"。
- `entity_name`：**重要的商业/办公/居住类实体名称，必须是完整的名称**。
- `Additional Information`：前面字段抽取完成后，剩余的全部文本；包括：市场内部区域名（如"江秾汇"）、楼层号、房间号、档口号、铺位号、单元号，以及括号内的补充说明中不属于实体名称的部分

注意事项：
  1. road_name不要提取市场/园区/小区内部的编号道路（如"八马路"、"5号路"、"东区一路"等）
  2. road_name不要提取非正式的道路俗称
  3. 如果地址中没有正式道路名称，road_name 留空
  4. 如果road_no前面紧挨着的内容不是 road_name，则留空
  5. road_no只提取主干路号，格式为：数字 + "号"
  6. road_no如果有"之X"、"之X至Y号"、"附X"、"支X"等后缀，只保留到第一个"之"、"附"、"支"之前
  7. road_no不要包含后面的单元号、楼层号、铺位号或其他任何内容
  8. entity_name不要提取道路类词（含"路、大道、马路、街、巷"等）
  9. entity_name不要提取市场内部的区域/通道编号（如"江秾汇"、"八马路"）
  10. entity_name不要提取单字母、单数字、纯编号
  11. entity_name不要提取未闭合括号的内容
  12. entity_name不要提取单独的实体后缀
  13. 当地址中同时出现公司全称和商业实体名称（如“XX服装城”“XX大厦”）时，优先提取商业实体名称。

输入：
{json.dumps(input_payload, ensure_ascii=False, indent=2)}

输出 JSON 结构：
{{
  "province": "",
  "city": "",
  "district": "",
  "town_street": "",
  "road_name": "",
  "road_no": "",
  "entity_name": "",
  "Additional Information": ""
}}
""".strip()


def extract_json_text(text):
    """从模型返回内容里提取 JSON 文本。"""
    cleaned_text = text.strip()

    if cleaned_text.startswith("```"):
        cleaned_text = re.sub(r"^```(?:json)?", "", cleaned_text).strip()
        cleaned_text = re.sub(r"```$", "", cleaned_text).strip()

    match = re.search(r"\{.*\}", cleaned_text, re.S)
    if match:
        return match.group(0)

    return cleaned_text


def truncate_text(text, max_length=500):
    """截断过长文本，避免错误信息过大。"""
    if len(text) <= max_length:
        return text
    return text[:max_length] + "..."


def call_deepseek(prompt):
    """调用硅基流动 DeepSeek API。"""
    if not DEEPSEEK_API_KEY:
        raise ValueError("请设置有效的 DEEPSEEK_API_KEY")

    last_error = None

    for attempt in range(1, API_RETRY_TIMES + 1):
        try:
            response = requests.post(
                DEEPSEEK_API_URL,
                headers={
                    "Authorization": f"Bearer {DEEPSEEK_API_KEY}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": DEEPSEEK_MODEL,
                    "messages": [
                        {
                            "role": "system",
                            "content": "你是一个严格的中文地址结构化助手，只返回 JSON。",
                        },
                        {
                            "role": "user",
                            "content": prompt,
                        },
                    ],
                    "temperature": 0.1,
                },
                timeout=120,
            )
            response.raise_for_status()

            response_data = response.json()
            if not isinstance(response_data, dict):
                raise ValueError(f"API响应不是字典: {truncate_text(str(response_data))}")

            choices = response_data.get("choices")
            if not isinstance(choices, list) or not choices:
                raise ValueError(f"API响应缺少choices: {truncate_text(json.dumps(response_data, ensure_ascii=False))}")

            first_choice = choices[0]
            if not isinstance(first_choice, dict):
                raise ValueError(f"choices[0]不是字典: {truncate_text(json.dumps(response_data, ensure_ascii=False))}")

            message = first_choice.get("message")
            if not isinstance(message, dict):
                raise ValueError(f"choices[0].message为空或不是字典: {truncate_text(json.dumps(response_data, ensure_ascii=False))}")

            content = message.get("content")
            if content is None:
                raise ValueError(f"choices[0].message.content为空: {truncate_text(json.dumps(response_data, ensure_ascii=False))}")

            if isinstance(content, str):
                return content

            raise ValueError(f"无法识别的message.content结构: {truncate_text(json.dumps(response_data, ensure_ascii=False))}")

        except Exception as error:
            last_error = error
            error_str = str(error)

            if "rate limit" in error_str.lower() or "quota" in error_str.lower():
                print(f"\n❌ API 配额或限流问题！错误: {error_str}")
                raise

            if attempt < API_RETRY_TIMES:
                print(f"  第 {attempt} 次调用失败，{API_RETRY_SLEEP_SECONDS}秒后重试... 错误: {error_str}")
                time.sleep(API_RETRY_SLEEP_SECONDS)

    raise last_error


def save_rows(header, rows):
    """将结果写入 Excel，输出列：company_address + 8个重构字段 + api_error + run_status。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"

    output_header = ["company_address"] + STRUCTURED_FIELDS + ["api_error", "run_status"]
    worksheet.append(output_header)

    for row in rows:
        addr = row.get("company_address", "")
        rebuilt_values = [row.get(f"rebuilt_{field_name}", "") for field_name in STRUCTURED_FIELDS]
        worksheet.append(
            [addr] + rebuilt_values + [row.get("api_error", ""), row.get("run_status", "")]
        )

    workbook.save(OUTPUT_PATH)
    print(f"✓ 已保存进度到 {OUTPUT_PATH.name}")


# ------------------------------------------------------------
# 信号处理函数
# ------------------------------------------------------------

def signal_handler(signum, frame):
    global interrupted
    interrupted = True
    print("\n\n" + "="*60)
    print("收到中断信号，正在保存当前进度...")
    print("="*60)

    if current_header is not None and current_rows is not None:
        try:
            save_rows(current_header, current_rows)
            print("✓ 进度已保存，可以安全退出")
        except Exception as e:
            print(f"✗ 保存失败: {e}")

    print("\n程序退出")
    sys.exit(0)


# ------------------------------------------------------------
# 核心逻辑
# ------------------------------------------------------------

def rebuild_row(row):
    """调用模型重构单行结果（无后处理，直接返回模型原始输出）。"""
    prompt = build_prompt(row)
    content = call_deepseek(prompt)
    json_text = extract_json_text(content)
    rebuilt_data = json.loads(json_text)
    if not isinstance(rebuilt_data, dict):
        raise ValueError(f"模型返回的JSON不是对象: {truncate_text(str(rebuilt_data))}")

    result = {}
    for field_name in STRUCTURED_FIELDS:
        result[field_name] = normalize_value(rebuilt_data.get(field_name, ""))

    return result


# ------------------------------------------------------------
# 主函数
# ------------------------------------------------------------

def main():
    global current_header, current_rows, interrupted

    start_time = time.time()

    signal.signal(signal.SIGINT, signal_handler)
    signal.signal(signal.SIGTERM, signal_handler)

    header, rows = load_rows()
    progress = load_existing_progress()
    merge_existing_progress(rows, progress)

    current_header = header
    current_rows = rows

    pending_rows = []
    skipped_addresses = []

    for row in rows:
        if is_row_successful(row):
            addr = row.get("company_address", "")
            skipped_addresses.append(addr)
        else:
            pending_rows.append(row)

    total_count = len(rows)
    pending_count = len(pending_rows)
    skipped_count = len(skipped_addresses)

    if skipped_count > 0:
        print(f"\n{'='*60}")
        print(f"跳过 {skipped_count} 条已成功处理的记录：")
        print(f"{'='*60}")
        for i, addr in enumerate(skipped_addresses[:20], 1):
            # 地址可能很长，截断显示
            short_addr = addr[:60] + "..." if len(addr) > 60 else addr
            print(f"  {i}. {short_addr}")
        if skipped_count > 20:
            print(f"  ... 还有 {skipped_count - 20} 条未显示")
        print(f"{'='*60}\n")

    print(f"总记录数: {total_count}")
    print(f"已处理(跳过): {skipped_count}")
    print(f"待处理: {pending_count}")
    print(f"保存间隔: 每 {SAVE_INTERVAL} 条")
    print(f"API 模型: {DEEPSEEK_MODEL}")
    print(f"提示: 按 Ctrl+C 可随时保存进度并退出\n")
    print(f"开始重构待处理记录...\n")

    for index, row in enumerate(pending_rows, start=1):
        if interrupted:
            break

        addr = row.get("company_address", "")
        short_addr = addr[:60] + "..." if len(addr) > 60 else addr
        print(f"当前处理进度: {index}/{pending_count} | 地址: {short_addr}")

        try:
            rebuilt_row = rebuild_row(row)
            for field_name in STRUCTURED_FIELDS:
                row[f"rebuilt_{field_name}"] = rebuilt_row[field_name]
            row["api_error"] = ""
            row["run_status"] = STATUS_SUCCESS
        except Exception as error:
            for field_name in STRUCTURED_FIELDS:
                row[f"rebuilt_{field_name}"] = ""
            row["api_error"] = str(error)
            row["run_status"] = STATUS_FAILED
            print(f"  ❌ 处理失败: {error}")

        if index % SAVE_INTERVAL == 0 or index == pending_count or interrupted:
            save_rows(header, rows)

    if not interrupted:
        save_rows(header, rows)

    print(f"\n{'='*60}")
    print(f"处理完成！")
    print(f"输入文件: {INPUT_PATH.name}")
    print(f"输出文件: {OUTPUT_PATH.name}")
    print(f"总记录数: {total_count}")
    processed_count = pending_count - (pending_count - index if interrupted else 0)
    print(f"本次处理: {processed_count}")
    print(f"跳过记录: {skipped_count}")
    print(f"{'='*60}")

    end_time = time.time()
    print(f"运行时间: {end_time - start_time:.4f}秒")


if __name__ == "__main__":
    main()